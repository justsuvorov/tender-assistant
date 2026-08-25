"""Сквозной прогон всех трёх этапов через оркестратор."""

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from tender_assistant.core.pydantic_models import DocumentStatus, FieldStatus
from tender_assistant.services.assistant import TenderAssistantService
from tests.fakes import Marker, ScriptedModel, as_json, fenced

pytestmark = pytest.mark.integration


@pytest.fixture
def result(api_request, scripted_model):
    return TenderAssistantService(request=api_request, ai_model=scripted_model).result()


class TestFullPipeline:
    def test_request_id_is_carried_through(self, result):
        assert result.request_id == 42

    def test_stage_one_produces_a_checked_document_list(self, result):
        assert [d.name for d in result.document_list.documents] == [
            "Выписка из ЕГРЮЛ",
            "Устав организации",
            "Бухгалтерский баланс за последний отчётный год",
            "Справка об отсутствии задолженности по налогам",
        ]
        assert [d.name for d in result.document_list.excluded] == [
            "Копия паспорта индивидуального предпринимателя"
        ]

    def test_stage_two_receives_the_checked_list(self, result):
        """Связь этапов: во второй уходит перечень после нормативной проверки."""
        assert [r.name for r in result.prepared_documents.rows] == [
            d.name for d in result.document_list.documents
        ]

    def test_stage_two_statuses(self, result):
        statuses = {r.name: r.status for r in result.prepared_documents.rows}

        assert statuses["Выписка из ЕГРЮЛ"] is DocumentStatus.FOUND
        assert statuses["Устав организации"] is DocumentStatus.FOUND
        assert statuses["Бухгалтерский баланс за последний отчётный год"] is DocumentStatus.CHECK
        assert statuses["Справка об отсутствии задолженности по налогам"] is DocumentStatus.MISSING

    def test_files_are_collected_into_one_folder(self, result, results_dir):
        komplekt = results_dir / "komplekt"

        assert {p.name for p in komplekt.iterdir()} == {
            "Выписка ЕГРЮЛ 2026-07-14.pdf",
            "Устав ООО Ромашка ред.5.docx",
            "Баланс_2025.xlsx",
        }

    def test_stage_three_fills_the_template(self, result):
        filled = Path(result.application.filled_path)
        assert filled.exists()

        cells = {r.cells[0].text: r.cells[1].text
                 for r in Document(str(filled)).tables[0].rows}
        assert cells["Полное наименование участника"] == (
            "Общество с ограниченной ответственностью «Ромашка»"
        )
        assert cells["ИНН"] == "7701234567"
        assert cells["Контактный телефон"] == "НЕ НАЙДЕНО"

    def test_paragraph_field_is_filled(self, result):
        text = "\n".join(
            p.text for p in Document(result.application.filled_path).paragraphs
        )
        assert "Руководитель организации: Иванов Иван Иванович" in text

    def test_colour_coding_reaches_the_document(self, result):
        rows = Document(result.application.filled_path).tables[0].rows
        highlights = [r.cells[1].paragraphs[0].runs[0].font.highlight_color
                      for r in rows]

        assert highlights[1] == WD_COLOR_INDEX.BRIGHT_GREEN   # ИНН — found
        assert highlights[2] == WD_COLOR_INDEX.YELLOW         # адрес — check
        assert highlights[3] == WD_COLOR_INDEX.RED            # телефон — missing

    def test_field_statuses(self, result):
        statuses = {f.label: f.status for f in result.application.fields}

        assert statuses["ИНН"] is FieldStatus.FOUND
        assert statuses["Юридический адрес"] is FieldStatus.CHECK
        assert statuses["Контактный телефон"] is FieldStatus.MISSING

    def test_every_stage_writes_a_report(self, result):
        for path in (
            result.document_list.report_path,
            result.prepared_documents.report_path,
            result.application.report_path,
        ):
            assert Path(path).exists()
            assert Path(path).read_text(encoding="utf-8").strip()

    def test_reports_land_in_the_requested_folder(self, result, results_dir):
        reports = sorted(p.name for p in results_dir.glob("*.md"))
        assert len(reports) == 3
        assert reports[0].startswith("application_fill_")
        assert reports[1].startswith("document_list_")
        assert reports[2].startswith("documents_status_")

    def test_source_documents_are_not_modified(
        self, api_request, scripted_model, requirements_docx, template_docx
    ):
        before = (requirements_docx.read_bytes(), template_docx.read_bytes())
        TenderAssistantService(request=api_request, ai_model=scripted_model).result()

        assert (requirements_docx.read_bytes(), template_docx.read_bytes()) == before


class TestPipelineDegradation:
    """Пайплайн обязан доходить до конца на неполных входных данных."""

    def test_no_normative_base(self, api_request, scripted_model):
        api_request.normative_base_folder = None
        result = TenderAssistantService(
            request=api_request, ai_model=scripted_model
        ).result()

        # Без нормативной базы перечень не фильтруется
        assert len(result.document_list.documents) == 5
        assert result.document_list.excluded == []
        assert Path(result.application.filled_path).exists()

    def test_no_knowledge_base_falls_back_to_normative(
        self, api_request, scripted_model
    ):
        api_request.knowledge_base_folder = None
        TenderAssistantService(request=api_request, ai_model=scripted_model).result()

        assert "юридическое лицо" in scripted_model.last_call(Marker.FIELD_VALUE)

    def test_empty_archive(self, api_request, scripted_model, tmp_path):
        empty = tmp_path / "пусто"
        empty.mkdir()
        api_request.documents_folder_path = str(empty)

        result = TenderAssistantService(
            request=api_request, ai_model=scripted_model
        ).result()

        assert all(r.status is DocumentStatus.MISSING
                   for r in result.prepared_documents.rows)
        assert Path(result.application.filled_path).exists()

    def test_no_documents_found_in_requirements(self, api_request, results_dir):
        """Раздела с документами нет — второй этап пуст, третий работает."""
        model = ScriptedModel({
            Marker.HEADERS: fenced({"headers": [], "confidence": "low"}),
            Marker.SECTION: as_json({"is_document_list": False, "documents": []}),
            Marker.FORM_FIELDS: as_json({"fields": [{"label": "ИНН"}]}),
            Marker.FIELD_VALUE: as_json(
                {"value": "7701234567", "status": "found", "source": "kb", "note": ""}
            ),
        })
        result = TenderAssistantService(request=api_request, ai_model=model).result()

        assert result.document_list.documents == []
        assert result.prepared_documents.rows == []
        assert result.application.fields[0].value == "7701234567"

    def test_unreadable_template_stops_the_run(self, api_request, scripted_model, tmp_path):
        """Неподдерживаемый формат шаблона — осмысленная ошибка, не падение."""
        bad = tmp_path / "form.rtf"
        bad.write_text("stub", encoding="utf-8")
        api_request.application_template_path = str(bad)

        with pytest.raises(ValueError, match="Unsupported file format"):
            TenderAssistantService(request=api_request, ai_model=scripted_model).result()


class TestModelUsage:
    def test_number_of_model_calls(self, api_request, scripted_model):
        TenderAssistantService(request=api_request, ai_model=scripted_model).result()

        assert scripted_model.call_count(Marker.HEADERS) == 1
        assert scripted_model.call_count(Marker.SECTION) == 1
        assert scripted_model.call_count(Marker.NORMATIVE) == 1
        assert scripted_model.call_count(Marker.FILE_MATCH) == 4   # по документу
        assert scripted_model.call_count(Marker.FORM_FIELDS) == 1
        assert scripted_model.call_count(Marker.FIELD_VALUE) == 5  # по полю

    def test_factory_is_used_when_no_model_given(self, api_request, monkeypatch, scripted_model):
        from tender_assistant.ai.model import ModelFactory

        monkeypatch.setattr(ModelFactory, "create", staticmethod(lambda: scripted_model))
        service = TenderAssistantService(request=api_request)

        assert service.ai_model is scripted_model
