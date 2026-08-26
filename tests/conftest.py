"""Общие фикстуры: синтетическая тендерная документация и фиктивная LLM.

Переменные окружения выставляются до импорта ``tender_assistant.core.config``,
потому что ``settings`` — синглтон, создаваемый на импорте модуля.
"""

import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.environ.setdefault("AI_PROVIDER", "gemini")
os.environ.setdefault("GEMINI_API_KEY", "test-key")
os.environ.setdefault("ANTHROPIC_API_KEY", "test-key")
os.environ.setdefault("PROMPT_CONTEXT_WINDOW", "0")

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from tests.fakes import Marker, ScriptedModel, as_json, fenced, value_between  # noqa: E402


# ── Исходные документы ────────────────────────────────────────────────────────

@pytest.fixture
def requirements_docx(tmp_path: Path) -> Path:
    """Требования тендера: разделы со стилевыми заголовками."""
    doc = Document()
    doc.add_heading("Конкурсная документация", level=1)

    doc.add_heading("1. Общие положения", level=2)
    doc.add_paragraph("Заказчик проводит открытый конкурс на оказание услуг.")

    doc.add_heading("2. Порядок подачи заявок", level=2)
    doc.add_paragraph("Заявки подаются в запечатанном конверте до 01.09.2026.")

    doc.add_heading("3. Состав заявки участника", level=2)
    doc.add_paragraph("3.1. Выписка из ЕГРЮЛ, выданная не ранее чем за 6 месяцев;")
    doc.add_paragraph("3.2. Устав организации в действующей редакции;")
    doc.add_paragraph("3.3. Бухгалтерский баланс за последний отчётный год;")
    doc.add_paragraph("3.4. Справка об отсутствии задолженности по налогам;")
    doc.add_paragraph("3.5. Копия паспорта индивидуального предпринимателя;")

    doc.add_heading("4. Критерии оценки", level=2)
    doc.add_paragraph("Оценка производится по цене и опыту участника.")

    path = tmp_path / "requirements.docx"
    doc.save(path)
    return path


@pytest.fixture
def template_docx(tmp_path: Path) -> Path:
    """Шаблон заявки: таблица реквизитов плюс строка с прочерком."""
    doc = Document()
    doc.add_heading("Заявка на участие в конкурсе", level=1)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    for i, label in enumerate([
        "Полное наименование участника",
        "ИНН",
        "Юридический адрес",
        "Контактный телефон",
    ]):
        table.rows[i].cells[0].text = label

    doc.add_paragraph("Руководитель организации: ______________")

    path = tmp_path / "form.docx"
    doc.save(path)
    return path


@pytest.fixture
def template_xlsx(tmp_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Наименование участника"
    sheet["A2"] = "ИНН"
    sheet["A3"] = "Контактный телефон"

    path = tmp_path / "form.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def archive_dir(tmp_path: Path) -> Path:
    """Архив документов организации, часть файлов — во вложенной папке."""
    archive = tmp_path / "archive"
    (archive / "uchreditelnye").mkdir(parents=True)

    (archive / "Выписка ЕГРЮЛ 2026-07-14.pdf").write_text("stub", encoding="utf-8")
    (archive / "Баланс_2025.xlsx").write_text("stub", encoding="utf-8")
    (archive / "Договор аренды офиса.docx").write_text("stub", encoding="utf-8")
    (archive / "uchreditelnye" / "Устав ООО Ромашка ред.5.docx").write_text(
        "stub", encoding="utf-8"
    )
    return archive


@pytest.fixture
def knowledge_dir(tmp_path: Path) -> Path:
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "company.md").write_text(
        "# Реквизиты организации\n\n"
        "Полное наименование: Общество с ограниченной ответственностью «Ромашка»\n"
        "ИНН: 7701234567\n"
        "Юридический адрес: 101000, г. Москва, ул. Мясницкая, д. 1\n"
        "Руководитель: Иванов Иван Иванович\n",
        encoding="utf-8",
    )
    return knowledge


@pytest.fixture
def normative_dir(tmp_path: Path) -> Path:
    normative = tmp_path / "normative"
    normative.mkdir()
    (normative / "rules.md").write_text(
        "Наша организация — юридическое лицо (ООО). "
        "Документы, требуемые от ИП и физических лиц, не предоставляются.\n",
        encoding="utf-8",
    )
    return normative


@pytest.fixture
def results_dir(tmp_path: Path) -> Path:
    results = tmp_path / "results"
    results.mkdir()
    return results


# ── Ответы фиктивной модели ───────────────────────────────────────────────────

DOCUMENTS_FROM_SECTION = [
    {"name": "Выписка из ЕГРЮЛ", "mandatory": True, "note": "не ранее 6 месяцев"},
    {"name": "Устав организации", "mandatory": True, "note": ""},
    {"name": "Бухгалтерский баланс за последний отчётный год", "mandatory": True, "note": ""},
    {"name": "Справка об отсутствии задолженности по налогам", "mandatory": True, "note": ""},
    {"name": "Копия паспорта индивидуального предпринимателя", "mandatory": True, "note": ""},
]

DOCUMENTS_AFTER_NORMATIVE = DOCUMENTS_FROM_SECTION[:4]

EXCLUDED_BY_NORMATIVE = [
    {"name": "Копия паспорта индивидуального предпринимателя",
     "note": "относится к ИП, наша организация — юридическое лицо"},
]

FILE_MATCHES = {
    "Выписка из ЕГРЮЛ": (["Выписка ЕГРЮЛ 2026-07-14.pdf"], "high"),
    "Устав организации": (["Устав ООО Ромашка ред.5.docx"], "high"),
    "Бухгалтерский баланс за последний отчётный год": (["Баланс_2025.xlsx"], "medium"),
}

FORM_FIELDS = [
    {"label": "Полное наименование участника", "anchor": "", "kind": "table_row"},
    {"label": "ИНН", "anchor": "", "kind": "table_row"},
    {"label": "Юридический адрес", "anchor": "", "kind": "table_row"},
    {"label": "Контактный телефон", "anchor": "", "kind": "table_row"},
    {"label": "Руководитель организации", "anchor": "", "kind": "line"},
]

FIELD_VALUES = {
    "Полное наименование участника": (
        "Общество с ограниченной ответственностью «Ромашка»", "found"),
    "ИНН": ("7701234567", "found"),
    "Юридический адрес": ("101000, г. Москва, ул. Мясницкая, д. 1", "check"),
    "Контактный телефон": ("", "missing"),
    "Руководитель организации": ("Иванов Иван Иванович", "found"),
}


def _match_file(query: str) -> str:
    target = value_between(query, Marker.FILE_MATCH)
    files, confidence = FILE_MATCHES.get(target, ([], "low"))
    return as_json({"files": files, "confidence": confidence, "note": ""})


def _fill_field(query: str) -> str:
    label = value_between(query, Marker.FIELD_VALUE)
    value, status = FIELD_VALUES.get(label, ("", "missing"))
    return as_json(
        {"value": value, "status": status, "source": "company.md", "note": ""}
    )


@pytest.fixture
def scripted_model() -> ScriptedModel:
    """Модель, отвечающая согласованно по всем трём этапам."""
    return ScriptedModel({
        Marker.HEADERS: fenced(
            {"headers": ["3. Состав заявки участника"], "confidence": "high"}
        ),
        Marker.SECTION: as_json(
            {"is_document_list": True, "documents": DOCUMENTS_FROM_SECTION}
        ),
        Marker.NORMATIVE: as_json(
            {"documents": DOCUMENTS_AFTER_NORMATIVE, "excluded": EXCLUDED_BY_NORMATIVE}
        ),
        Marker.FILE_MATCH: _match_file,
        Marker.FORM_FIELDS: as_json({"fields": FORM_FIELDS}),
        Marker.FIELD_VALUE: _fill_field,
    })


@pytest.fixture
def api_request(
    requirements_docx, archive_dir, results_dir, template_docx,
    normative_dir, knowledge_dir,
):
    from tender_assistant.models.request import APIRequest

    return APIRequest(
        message_id=42,
        file_path=str(requirements_docx),
        documents_folder_path=str(archive_dir),
        results_path=str(results_dir),
        application_template_path=str(template_docx),
        normative_base_folder=str(normative_dir),
        knowledge_base_folder=str(knowledge_dir),
        result_folder_name="komplekt",
    )
