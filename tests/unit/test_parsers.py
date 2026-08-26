"""Чтение документов и разбор их структуры."""

import pytest
from docx import Document
from openpyxl import Workbook

from tender_assistant.core.parsers import (
    DataParser,
    MarkdownOutline,
    MarkdownTableBuilder,
    Section,
    find_inline_blanks,
    render_blanks_context,
    xml_path,
)


class TestMarkdownTableBuilder:
    def test_normalise_pads_short_rows(self):
        rows = MarkdownTableBuilder.normalise([["а", "б", "в"], ["г"]])
        assert rows == [["а", "б", "в"], ["г", "", ""]]

    def test_normalise_stringifies_and_strips(self):
        rows = MarkdownTableBuilder.normalise([[1, None, " x \n y "]])
        assert rows == [["1", "", "x   y"]]

    def test_from_rows_builds_header_and_separator(self):
        md = MarkdownTableBuilder.from_rows([["Поле", "Значение"], ["ИНН", "770"]])
        lines = md.splitlines()
        assert lines[0].startswith("| Поле")
        assert set(lines[1]) <= set("| -")
        assert "ИНН" in lines[2]

    def test_empty_input(self):
        assert MarkdownTableBuilder.from_rows([]) == ""
        assert MarkdownTableBuilder.normalise([]) == []


class TestDataParser:
    def test_reads_word_headings_as_markdown(self, requirements_docx):
        markdown = DataParser(str(requirements_docx)).origin_data()
        assert "# Конкурсная документация" in markdown
        assert "## 3. Состав заявки участника" in markdown
        assert "Выписка из ЕГРЮЛ" in markdown

    def test_reads_word_tables(self, template_docx):
        markdown = DataParser(str(template_docx)).origin_data()
        assert "| Полное наименование участника" in markdown

    def test_reads_excel_sheets(self, template_xlsx):
        markdown = DataParser(str(template_xlsx)).origin_data()
        assert "## Лист:" in markdown
        assert "ИНН" in markdown

    def test_unsupported_format_raises(self, tmp_path):
        path = tmp_path / "notes.rtf"
        path.write_text("stub", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported file format"):
            DataParser(str(path))

    def test_missing_file_raises_a_clear_error(self, tmp_path):
        """Иначе наружу летит PackageNotFoundError из python-docx."""
        with pytest.raises(FileNotFoundError, match="Файл не найден"):
            DataParser(str(tmp_path / "нет-такого.docx"))

    def test_directory_instead_of_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="не является файлом"):
            DataParser(str(tmp_path))

    @pytest.mark.parametrize("path", ["", None])
    def test_empty_path_raises(self, path):
        with pytest.raises(FileNotFoundError, match="Не указан путь"):
            DataParser(path)

    def test_excel_skips_empty_sheets(self, tmp_path):
        workbook = Workbook()
        workbook.active["A1"] = "Данные"
        workbook.create_sheet("Пустой")
        path = tmp_path / "book.xlsx"
        workbook.save(path)

        assert "Пустой" not in DataParser(str(path)).origin_data()

    def test_blank_lines_are_collapsed(self, tmp_path):
        doc = Document()
        doc.add_paragraph("Первый")
        for _ in range(5):
            doc.add_paragraph("")
        doc.add_paragraph("Второй")
        path = tmp_path / "sparse.docx"
        doc.save(path)

        assert "\n\n\n" not in DataParser(str(path)).origin_data()


class TestMarkdownOutlineWithHeadings:
    @pytest.fixture
    def outline(self):
        return MarkdownOutline(
            "# Раздел А\nтекст А\n\n## Раздел Б\nтекст Б\n\n### Раздел В\nтекст В"
        )

    def test_titles_and_levels(self, outline):
        assert outline.titles() == ["Раздел А", "Раздел Б", "Раздел В"]
        assert [s.level for s in outline.sections] == [1, 2, 3]

    def test_body_belongs_to_its_heading(self, outline):
        assert outline.find("Раздел Б").body == "текст Б"

    def test_last_section_body_is_captured(self, outline):
        assert outline.find("Раздел В").body == "текст В"

    def test_section_text_includes_heading(self, outline):
        assert outline.find("Раздел Б").text == "Раздел Б\n\nтекст Б"

    def test_as_list_is_numbered(self, outline):
        assert outline.as_list().splitlines()[1] == "2. Раздел Б"

    def test_text_before_first_heading_is_ignored(self):
        outline = MarkdownOutline("преамбула\n# Раздел\nтело")
        assert outline.titles() == ["Раздел"]
        assert outline.find("Раздел").body == "тело"

    def test_trailing_hashes_are_stripped(self):
        assert MarkdownOutline("## Раздел ##\nтело").titles() == ["Раздел"]


class TestMarkdownOutlineFind:
    @pytest.fixture
    def outline(self):
        return MarkdownOutline("## 3. Состав заявки участника\nтело\n\n## 4. Оценка\n-")

    def test_exact_match(self, outline):
        assert outline.find("3. Состав заявки участника").level == 2

    def test_case_and_punctuation_insensitive(self, outline):
        assert outline.find("состав заявки участника") is not None

    def test_partial_match(self, outline):
        """Модель часто возвращает заголовок без нумерации."""
        assert outline.find("Состав заявки").title == "3. Состав заявки участника"

    def test_unknown_heading(self, outline):
        assert outline.find("Раздела нет") is None

    @pytest.mark.parametrize("title", ["", None])
    def test_empty_query(self, outline, title):
        assert outline.find(title) is None


class TestMarkdownOutlineHeuristics:
    """Документы без стилевых заголовков — сплошь и рядом в тендерах."""

    @pytest.fixture
    def outline(self):
        return MarkdownOutline(
            "Извещение о проведении конкурса.\n"
            "1. Общие положения\n"
            "Заказчик проводит конкурс.\n"
            "2.1. Состав заявки\n"
            "Устав; выписка из ЕГРЮЛ;\n"
            "ПРИЛОЖЕНИЕ N 1\n"
            "форма заявки\n"
        )

    def test_numbered_paragraphs_become_headings(self, outline):
        assert outline.titles() == [
            "1. Общие положения", "2.1. Состав заявки", "ПРИЛОЖЕНИЕ N 1"
        ]

    def test_nesting_level_from_numbering(self, outline):
        assert [s.level for s in outline.sections] == [1, 2, 1]

    def test_body_is_captured(self, outline):
        assert outline.find("2.1. Состав заявки").body == "Устав; выписка из ЕГРЮЛ;"

    def test_enumeration_items_are_not_headings(self):
        """Пункт перечня заканчивается «;» — это тело раздела, а не заголовок."""
        outline = MarkdownOutline("1. Требования\n3.1. Устав организации;\n")
        assert outline.titles() == ["1. Требования"]

    def test_sentences_are_not_headings(self):
        outline = MarkdownOutline(
            "1. Требования\n2. Заявка подаётся в срок. Конверт запечатывается.\n"
        )
        assert outline.titles() == ["1. Требования"]

    def test_table_rows_are_not_headings(self):
        outline = MarkdownOutline("1. Требования\n| 2. Ячейка таблицы |\n")
        assert outline.titles() == ["1. Требования"]

    def test_heuristics_yield_to_real_headings(self):
        """При наличии стилевых заголовков нумерованные абзацы не трогаем."""
        outline = MarkdownOutline("# Настоящий заголовок\n1. Не заголовок\nтело")
        assert outline.titles() == ["Настоящий заголовок"]

    def test_document_without_any_structure(self):
        outline = MarkdownOutline("Просто сплошной текст без разделов.")
        assert not outline.has_headings
        assert outline.titles() == []

    def test_empty_document(self):
        outline = MarkdownOutline("")
        assert not outline.has_headings
        assert outline.as_list() == ""


class TestSection:
    def test_text_without_body(self):
        assert Section("Заголовок", 1).text == "Заголовок"

    def test_repr_shows_title(self):
        assert "Заголовок" in repr(Section("Заголовок", 2))


class TestXmlPath:
    """id() Python-обёрток python-docx/lxml ненадёжен как ключ идентичности —
    объекты эфемерны, GC может отдать тот же id() совсем другому элементу
    на настоящем (не игрушечном) документе. xml_path строит ключ из позиции
    в дереве, а не из адреса Python-объекта."""

    def test_same_element_same_path(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        assert xml_path(cell._tc) == xml_path(table.rows[0].cells[0]._tc)

    def test_different_elements_different_paths(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        assert xml_path(table.rows[0].cells[0]._tc) != xml_path(table.rows[1].cells[0]._tc)


def _add_cell_text(cell, text: str) -> None:
    cell.paragraphs[0].add_run(text)


class TestFindInlineBlanks:
    """Пропуски внутри абзацев ячеек таблиц — регрессия на реальную форму
    Росатова (223-ФЗ): пункт вида «15.1» — три РАЗНЫХ по смыслу пропуска
    в одном предложении, каждому нужно своё значение."""

    def test_two_blanks_with_hints_in_one_paragraph_are_found(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        _add_cell_text(
            table.rows[0].cells[0],
            "_____ (наименование участника закупки) зарегистрирован в "
            "_____ (наименование государства) в установленном порядке.",
        )

        blanks = find_inline_blanks(doc)

        assert len(blanks) == 2
        assert blanks[0].label == "наименование участника закупки"
        assert blanks[1].label == "наименование государства"

    def test_single_blank_in_a_cell_is_not_treated_as_inline(self):
        """Один пропуск на абзац — это зона старого механизма (метка в
        ячейке → значение в соседней), не инлайн-детектора. Порог "2 и
        больше" — намеренный, чтобы не дублировать и не конфликтовать."""
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Руководитель"
        _add_cell_text(table.rows[0].cells[1], "______________")

        assert find_inline_blanks(doc) == []

    def test_body_paragraphs_are_not_scanned(self):
        """Инлайн-детектор — только ячейки таблиц. Одиночные и множественные
        пропуски в body-абзацах уже работают через старый механизм
        (WordApplicationWriter._fill_in_paragraphs); дублировать эту зону
        не нужно — риск двойной записи в один абзац."""
        doc = Document()
        doc.add_paragraph(
            "___ (участник) и ___ (страна) — оба пропуска вне таблицы."
        )
        assert find_inline_blanks(doc) == []

    def test_fallback_label_when_no_hint(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        _add_cell_text(
            table.rows[0].cells[0],
            "Настоящим подтверждаем: ___ является учредителем ___ полностью.",
        )

        blanks = find_inline_blanks(doc)
        assert len(blanks) == 2
        assert blanks[0].label != ""
        assert "поясняющей подписи" in blanks[1].label or blanks[1].label

    def test_ids_are_stable_and_sequential(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        _add_cell_text(table.rows[0].cells[0], "___ (а) и ___ (б) и ___ (в).")

        blanks = find_inline_blanks(doc)
        assert [b.id for b in blanks] == [1, 2, 3]

    def test_merged_cell_is_scanned_once_not_per_row(self):
        """Дедупликация через xml_path: без неё абзацы объединённой ячейки
        попали бы в результат по разу на каждую охваченную строку."""
        doc = Document()
        table = doc.add_table(rows=3, cols=1)
        table.cell(0, 0).merge(table.cell(2, 0))
        _add_cell_text(
            table.rows[0].cells[0],
            "___ (участник) подтверждает ___ (условие) в полном объёме.",
        )

        blanks = find_inline_blanks(doc)
        assert len(blanks) == 2  # не 6 (2 пропуска × 3 строки объединения)

    def test_many_unrelated_cells_do_not_hide_the_match(self):
        """Регрессионный сценарий: на реальном документе с десятками ячеек
        id() python-docx обёрток совпадал у совершенно разных ячеек из-за
        переиспользования GC — из-за этого искомый абзац считался
        "уже виденным" и пропускался. Таблица с большим числом строк — сеть
        для похожего класса ошибок, даже не гарантируя точное воспроизведение."""
        doc = Document()
        table = doc.add_table(rows=30, cols=2)
        for i in range(30):
            table.rows[i].cells[0].text = f"Строка {i}"
            table.rows[i].cells[1].text = "просто текст без пропусков"

        _add_cell_text(
            table.rows[29].cells[1],
            "___ (участник) и ___ (страна) — в последней строке таблицы.",
        )

        blanks = find_inline_blanks(doc)
        assert len(blanks) == 2

    def test_nested_table_is_scanned(self):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        inner = outer.rows[0].cells[0].add_table(rows=1, cols=1)
        _add_cell_text(
            inner.rows[0].cells[0], "___ (участник) и ___ (страна) во вложенной таблице."
        )

        blanks = find_inline_blanks(doc)
        assert len(blanks) == 2

    def test_no_tables_at_all(self):
        doc = Document()
        doc.add_paragraph("Просто текст без таблиц.")
        assert find_inline_blanks(doc) == []


class TestRenderBlanksContext:
    def test_markers_replace_blank_spans(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        _add_cell_text(
            table.rows[0].cells[0],
            "___ (участник) зарегистрирован в ___ (страна) в установленном порядке.",
        )
        blanks = find_inline_blanks(doc)

        rendered = render_blanks_context(blanks)

        assert "[[1]]" in rendered and "[[2]]" in rendered
        assert "___" not in rendered
        assert "(участник)" in rendered  # поясняющая скобка остаётся как контекст

    def test_multiple_paragraphs_are_joined(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        _add_cell_text(table.rows[0].cells[0], "___ (а) и ___ (б) в первом.")
        _add_cell_text(table.rows[1].cells[0], "___ (в) и ___ (г) во втором.")
        blanks = find_inline_blanks(doc)

        rendered = render_blanks_context(blanks)

        assert rendered.count("[[") == 4
        assert "\n\n" in rendered  # абзацы разделены

    def test_empty_list(self):
        assert render_blanks_context([]) == ""
