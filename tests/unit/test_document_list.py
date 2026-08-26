"""Этап 1: поиск раздела с перечнем документов и его разбор."""

import pytest

from tender_assistant.ai.postprocessor import DocumentListResponse, SectionsMatcherResponse
from tender_assistant.core.parsers import DataParser
from tender_assistant.documents.document_list import (
    ContextMatcher,
    DocumentList,
    DocumentListExtractor,
    DocumentListOutput,
    DocumentSections,
    NormativeChecker,
    SectionsMatcher,
)
from tender_assistant.reports.report_export import DocumentListReport
from tests.fakes import BrokenModel, Marker, ScriptedModel, as_json, fenced


@pytest.fixture
def sections(requirements_docx):
    return DocumentSections(data_parser=DataParser(file_path=str(requirements_docx)))


class TestDocumentSections:
    def test_markdown_is_read_once(self, sections):
        """Документ разбирается лениво и кешируется: парсинг дорогой."""
        assert sections.markdown is sections.markdown

    def test_headings_are_detected(self, sections):
        assert sections.has_headings()

    def test_sections_list_is_numbered(self, sections):
        listing = sections.sections_list()
        assert "1. Конкурсная документация" in listing
        assert "4. 3. Состав заявки участника" in listing

    def test_section_text_returns_only_that_section(self, sections):
        text = sections.section_text("3. Состав заявки участника")
        assert "Выписка из ЕГРЮЛ" in text
        assert "Критерии оценки" not in text

    def test_unknown_header_falls_back_to_whole_document(self, sections):
        """Заголовок не найден — лучше отдать модели весь документ, чем ничего."""
        text = sections.section_text("Такого раздела нет")
        assert text == sections.markdown

    def test_document_without_headings(self, tmp_path):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Сплошной текст требований без разделов.")
        path = tmp_path / "flat.docx"
        doc.save(path)

        flat = DocumentSections(data_parser=DataParser(file_path=str(path)))
        assert not flat.has_headings()
        assert flat.sections_list() == ""


class TestSectionsMatcher:
    def test_returns_headers_from_model(self):
        model = ScriptedModel({Marker.HEADERS: fenced({"headers": ["Состав заявки"]})})
        matcher = SectionsMatcher(ai_model=model)
        assert matcher.result("1. Состав заявки") == ["Состав заявки"]

    def test_headings_are_passed_into_the_prompt(self):
        model = ScriptedModel({Marker.HEADERS: fenced({"headers": []})})
        SectionsMatcher(ai_model=model).result("1. Состав заявки")
        assert "1. Состав заявки" in model.last_call()

    def test_empty_headings_skip_the_model(self):
        """Без заголовков спрашивать модель не о чем — экономим вызов."""
        model = ScriptedModel({})
        assert SectionsMatcher(ai_model=model).result("   ") == []
        assert model.calls == []

    def test_garbage_answer_becomes_a_pseudo_header(self):
        """Текстовый откат превращает мусор в «заголовок» — он не найдётся
        в документе, и разбор уйдёт на весь текст (см. тест ниже)."""
        matcher = SectionsMatcher(ai_model=BrokenModel("не понимаю"))
        assert matcher.result("1. Состав заявки") == ["не понимаю"]

    def test_hallucinated_header_leads_to_whole_document(self, sections, results_dir):
        model = ScriptedModel({
            Marker.HEADERS: "выдуманный заголовок",
            Marker.SECTION: as_json(
                {"is_document_list": True, "documents": [{"name": "Устав"}]}
            ),
        })
        result = DocumentList(
            extractor=DocumentListExtractor(
                document_sections=sections,
                sections_matcher=SectionsMatcher(ai_model=model),
                context_matcher=ContextMatcher(ai_model=model),
            ),
            output=DocumentListOutput(results_path=str(results_dir)),
            report=DocumentListReport(output_dir=str(results_dir)),
        ).result()

        assert [d.name for d in result.documents] == ["Устав"]
        assert "Критерии оценки" in model.last_call(Marker.SECTION)


class TestContextMatcher:
    def test_extracts_documents(self):
        model = ScriptedModel({Marker.SECTION: as_json(
            {"is_document_list": True, "documents": [{"name": "Устав"}]}
        )})
        assert ContextMatcher(ai_model=model).result("текст")[0]["name"] == "Устав"

    def test_section_without_documents(self):
        model = ScriptedModel({Marker.SECTION: as_json(
            {"is_document_list": False, "documents": []}
        )})
        assert ContextMatcher(ai_model=model).result("текст") == []

    def test_empty_section_skips_the_model(self):
        model = ScriptedModel({})
        assert ContextMatcher(ai_model=model).result("  ") == []
        assert model.calls == []


class TestNormativeChecker:
    DOCS = [{"name": "Устав", "mandatory": True, "note": ""},
            {"name": "Паспорт ИП", "mandatory": True, "note": ""}]

    def test_filters_by_normative_base(self, normative_dir):
        model = ScriptedModel({Marker.NORMATIVE: as_json({
            "documents": [{"name": "Устав"}],
            "excluded": [{"name": "Паспорт ИП", "note": "мы юрлицо"}],
        })})
        checker = NormativeChecker(
            ai_model=model, normative_base_folder=str(normative_dir)
        )
        result = checker.result(self.DOCS)

        assert [d["name"] for d in result["documents"]] == ["Устав"]
        assert [d["name"] for d in result["excluded"]] == ["Паспорт ИП"]

    def test_normative_text_reaches_the_prompt(self, normative_dir):
        model = ScriptedModel({Marker.NORMATIVE: as_json(
            {"documents": [{"name": "Устав"}], "excluded": []}
        )})
        NormativeChecker(
            ai_model=model, normative_base_folder=str(normative_dir)
        ).result(self.DOCS)

        assert "юридическое лицо" in model.last_call()

    @pytest.mark.parametrize("folder", [None, ""])
    def test_without_base_the_list_passes_through(self, folder):
        model = ScriptedModel({})
        checker = NormativeChecker(ai_model=model, normative_base_folder=folder)

        assert checker.result(self.DOCS)["documents"] == self.DOCS
        assert model.calls == []

    def test_empty_list_skips_the_model(self, normative_dir):
        model = ScriptedModel({})
        checker = NormativeChecker(
            ai_model=model, normative_base_folder=str(normative_dir)
        )
        assert checker.result([]) == {"documents": [], "excluded": []}
        assert model.calls == []

    def test_model_dropping_everything_is_overruled(self, normative_dir):
        """Пустой результат фильтрации — почти всегда ошибка модели."""
        model = ScriptedModel({Marker.NORMATIVE: as_json(
            {"documents": [], "excluded": [{"name": "Устав"}]}
        )})
        checker = NormativeChecker(
            ai_model=model, normative_base_folder=str(normative_dir)
        )
        result = checker.result(self.DOCS)

        assert result["documents"] == self.DOCS
        assert result["excluded"] == []


class TestDocumentList:
    """Оркестратор: извлечь → проверить по норме → сохранить.

    Чтение требований и запись файлов живут в DocumentListExtractor
    и DocumentListOutput.
    """

    def _build(self, sections, model, results_dir, normative_folder=None):
        return DocumentList(
            extractor=DocumentListExtractor(
                document_sections=sections,
                sections_matcher=SectionsMatcher(
                    ai_model=model, response_post_processor=SectionsMatcherResponse()
                ),
                context_matcher=ContextMatcher(
                    ai_model=model, response_post_processor=DocumentListResponse()
                ),
            ),
            normative_checker=NormativeChecker(
                ai_model=model, normative_base_folder=normative_folder
            ) if normative_folder else None,
            output=DocumentListOutput(results_path=str(results_dir)),
            report=DocumentListReport(output_dir=str(results_dir)),
        )

    def test_takes_at_most_five_constructor_parameters(self):
        """Оркестратор не должен обрастать параметрами: всё, что сверх
        композиции зависимостей, — признак утёкшей в него работы."""
        import inspect

        parameters = inspect.signature(DocumentList.__init__).parameters
        assert len(parameters) - 1 <= 5  # без self

    def test_happy_path(self, sections, scripted_model, results_dir, normative_dir):
        result = self._build(
            sections, scripted_model, results_dir, str(normative_dir)
        ).result()

        assert [d.name for d in result.documents] == [
            "Выписка из ЕГРЮЛ",
            "Устав организации",
            "Бухгалтерский баланс за последний отчётный год",
            "Справка об отсутствии задолженности по налогам",
        ]
        assert [d.name for d in result.excluded] == [
            "Копия паспорта индивидуального предпринимателя"
        ]
        assert result.source_headers == ["3. Состав заявки участника"]

    def test_only_the_matched_section_is_read(
        self, sections, scripted_model, results_dir
    ):
        self._build(sections, scripted_model, results_dir).result()

        section_prompt = scripted_model.last_call(Marker.SECTION)
        assert "Выписка из ЕГРЮЛ" in section_prompt
        assert "Критерии оценки" not in section_prompt

    def test_report_is_written(self, sections, scripted_model, results_dir, normative_dir):
        from pathlib import Path

        result = self._build(
            sections, scripted_model, results_dir, str(normative_dir)
        ).result()

        report = Path(result.report_path).read_text(encoding="utf-8")
        assert "Перечень документов для заявки" in report
        assert "Выписка из ЕГРЮЛ" in report
        assert "Исключено при проверке по нормативной базе" in report

    def test_formatted_document_is_written_alongside_the_report(
        self, sections, scripted_model, results_dir, normative_dir
    ):
        """Перечень оформляется в Word вдобавок к markdown-отчёту, не вместо."""
        from docx import Document as ReadDocx
        from pathlib import Path

        result = self._build(
            sections, scripted_model, results_dir, str(normative_dir)
        ).result()

        assert result.formatted_path is not None
        assert result.formatted_path != result.report_path
        assert Path(result.formatted_path).exists()
        assert Path(result.formatted_path).suffix == ".docx"

        table = ReadDocx(result.formatted_path).tables[0]
        assert [row.cells[1].text for row in table.rows[1:]] == [
            d.name for d in result.documents
        ]

    def test_without_normative_checker_nothing_is_excluded(
        self, sections, scripted_model, results_dir
    ):
        result = self._build(sections, scripted_model, results_dir).result()
        assert len(result.documents) == 5
        assert result.excluded == []

    def test_document_without_headings_is_read_whole(
        self, tmp_path, results_dir
    ):
        """Ветка схемы «заголовков нет» — раздел не ищется вовсе."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("К заявке прилагаются: устав; выписка из ЕГРЮЛ.")
        path = tmp_path / "flat.docx"
        doc.save(path)

        model = ScriptedModel({Marker.SECTION: as_json(
            {"is_document_list": True, "documents": [{"name": "Устав"}]}
        )})
        flat = DocumentSections(data_parser=DataParser(file_path=str(path)))
        result = self._build(flat, model, results_dir).result()

        assert [d.name for d in result.documents] == ["Устав"]
        assert result.source_headers == []
        assert model.call_count(Marker.HEADERS) == 0

    def test_section_without_documents_falls_back_to_whole_document(
        self, sections, results_dir
    ):
        """Заголовок нашли, но перечня в разделе нет — читаем документ целиком."""
        calls = {"n": 0}

        def section_reply(query: str) -> str:
            calls["n"] += 1
            if calls["n"] == 1:
                return as_json({"is_document_list": False, "documents": []})
            return as_json(
                {"is_document_list": True, "documents": [{"name": "Устав"}]}
            )

        model = ScriptedModel({
            Marker.HEADERS: fenced({"headers": ["4. Критерии оценки"]}),
            Marker.SECTION: section_reply,
        })
        result = self._build(sections, model, results_dir).result()

        assert [d.name for d in result.documents] == ["Устав"]
        assert calls["n"] == 2
        assert "Конкурсная документация" in model.last_call(Marker.SECTION)

    def test_documents_from_several_sections_are_deduplicated(
        self, sections, results_dir
    ):
        model = ScriptedModel({
            Marker.HEADERS: fenced({"headers": [
                "3. Состав заявки участника", "4. Критерии оценки",
            ]}),
            Marker.SECTION: as_json({
                "is_document_list": True,
                "documents": [{"name": "Устав"}, {"name": "УСТАВ"}],
            }),
        })
        result = self._build(sections, model, results_dir).result()

        assert [d.name for d in result.documents] == ["Устав"]
        assert model.call_count(Marker.SECTION) == 2
