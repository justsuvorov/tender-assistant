"""Этап 3: разбор шаблона заявки и заполнение полей."""

from pathlib import Path

import pytest

from docx import Document

from tender_assistant.application.application import (
    AITenderForm,
    ApplicationOutput,
    CompositeTenderForm,
    FormTemplate,
    InlineBlanksForm,
    InlineBlanksQuery,
    KnowledgeBaseContext,
    TenderAIQuery,
    TenderApplication,
)
from tender_assistant.core.parsers import find_inline_blanks
from tender_assistant.core.pydantic_models import FieldStatus, FilledField
from tender_assistant.reports.report_export import TenderApplicationReport
from tests.fakes import BrokenModel, CachingScriptedModel, Marker, ScriptedModel, as_json


class TestTenderAIQuery:
    def test_returns_value_and_status(self):
        model = ScriptedModel({Marker.FIELD_VALUE: as_json(
            {"value": "7701234567", "status": "found", "source": "kb", "note": ""}
        )})
        result = TenderAIQuery(ai_model=model).result("ИНН", "контекст")

        assert result["value"] == "7701234567"
        assert result["status"] == "found"

    def test_field_and_context_reach_the_prompt(self):
        model = ScriptedModel({Marker.FIELD_VALUE: as_json({"value": "", "status": "missing"})})
        TenderAIQuery(ai_model=model).result("ИНН", "ИНН: 7701234567")

        prompt = model.last_call()
        assert "ИНН: 7701234567" in prompt
        assert "### ЗАПРАШИВАЕМОЕ ПОЛЕ" in prompt

    def test_garbage_answer_requires_a_check(self):
        result = TenderAIQuery(ai_model=BrokenModel("7701234567")).result("ИНН", "к")
        assert result["status"] == "check"


class TestTenderAIQueryPromptCaching:
    """Регрессия: живой прогон на ~200К-символьном документе тендера показал
    ~82К токенов контекста на КАЖДЫЙ из 47 запросов по полям заявки — материалы
    отправлялись целиком заново на каждое поле, не переиспользуясь. Провайдеры
    с prompt caching (Anthropic) должны получать материалы отдельным,
    неизменным между полями кэшируемым блоком."""

    def test_caching_model_receives_materials_as_a_stable_prefix(self):
        model = CachingScriptedModel({Marker.FIELD_VALUE: as_json(
            {"value": "7701234567", "status": "found", "source": "kb", "note": ""}
        )})
        TenderAIQuery(ai_model=model).result("ИНН", "ИНН: 7701234567 в базе знаний")

        assert len(model.cache_calls) == 1
        cache_prefix, query = model.cache_calls[0]
        assert "ИНН: 7701234567 в базе знаний" in cache_prefix
        assert "ИНН" in query
        assert "### МАТЕРИАЛЫ" in cache_prefix
        assert "### ЗАПРАШИВАЕМОЕ ПОЛЕ" in query

    def test_materials_are_identical_across_fields(self):
        """Смысл кэширования: один и тот же префикс на разные поля одной
        заявки — иначе Anthropic не сможет переиспользовать кэш."""
        model = CachingScriptedModel({Marker.FIELD_VALUE: as_json(
            {"value": "X", "status": "found", "source": "", "note": ""}
        )})
        query_maker = TenderAIQuery(ai_model=model)
        context = "общие материалы заявки"

        query_maker.result("ИНН", context)
        query_maker.result("Адрес", context)

        prefixes = [prefix for prefix, _ in model.cache_calls]
        assert prefixes[0] == prefixes[1]

    def test_materials_are_not_trimmed_for_a_caching_model(self):
        """Обрезка под конкретное поле сломала бы совпадение префиксов —
        для кэширующих моделей материалы уходят целиком."""
        big_context = "\n\n".join(f"# Раздел {i}\nпосторонний текст" for i in range(50))
        model = CachingScriptedModel({Marker.FIELD_VALUE: as_json(
            {"value": "X", "status": "found", "source": "", "note": ""}
        )})
        TenderAIQuery(ai_model=model).result("ИНН", big_context)

        cache_prefix, _ = model.cache_calls[0]
        assert cache_prefix.count("посторонний текст") == 50

    def test_non_caching_model_gets_a_single_concatenated_query(self):
        """Без поддержки кэширования — как раньше: один текст, без разбивки
        на префикс и запрос."""
        model = ScriptedModel({Marker.FIELD_VALUE: as_json(
            {"value": "X", "status": "found", "source": "", "note": ""}
        )})
        TenderAIQuery(ai_model=model).result("ИНН", "материалы")

        assert model.call_count(Marker.FIELD_VALUE) == 1
        assert "материалы" in model.last_call()


def _multi_blank_docx(path) -> None:
    """Строка вида «15.1» — несколько разных по смыслу пропусков в одном
    абзаце ячейки, как в реальной форме 223-ФЗ. Вторая колонка —
    «Предложение участника», туда пишутся подобранные значения."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run(
        "___ (наименование участника закупки) зарегистрирован в "
        "___ (наименование государства) в установленном порядке."
    )
    table.rows[0].cells[1].text = "«Да» / «Нет»"
    doc.save(path)


class TestInlineBlanksQuery:
    """Пакетный запрос по нескольким пропускам одного абзаца — один вызов
    модели на ВСЕ пропуски шаблона, а не по вызову на каждый."""

    def test_batches_all_blanks_into_one_call(self, tmp_path):
        path = tmp_path / "form.docx"
        _multi_blank_docx(path)
        blanks = find_inline_blanks(Document(str(path)))

        model = CachingScriptedModel({Marker.INLINE_BLANKS: as_json({
            "1": {"value": "САО «ВСК»", "status": "found", "source": "", "note": ""},
            "2": {"value": "Россия", "status": "found", "source": "", "note": ""},
        })})
        answers = InlineBlanksQuery(ai_model=model).result(blanks, "материалы")

        assert len(model.cache_calls) == 1
        assert answers["1"]["value"] == "САО «ВСК»"
        assert answers["2"]["value"] == "Россия"

    def test_materials_go_into_the_cache_prefix(self, tmp_path):
        path = tmp_path / "form.docx"
        _multi_blank_docx(path)
        blanks = find_inline_blanks(Document(str(path)))

        model = CachingScriptedModel({Marker.INLINE_BLANKS: as_json({
            "1": {"value": "X", "status": "found", "source": "", "note": ""},
            "2": {"value": "Y", "status": "found", "source": "", "note": ""},
        })})
        InlineBlanksQuery(ai_model=model).result(blanks, "уникальные материалы")

        cache_prefix, query = model.cache_calls[0]
        assert "уникальные материалы" in cache_prefix
        assert "[[1]]" in query and "[[2]]" in query

    def test_no_blanks_skips_the_model(self):
        model = CachingScriptedModel({})
        assert InlineBlanksQuery(ai_model=model).result([], "материалы") == {}
        assert model.cache_calls == []

    def test_non_caching_model_gets_one_concatenated_query(self, tmp_path):
        path = tmp_path / "form.docx"
        _multi_blank_docx(path)
        blanks = find_inline_blanks(Document(str(path)))

        model = ScriptedModel({Marker.INLINE_BLANKS: as_json({
            "1": {"value": "X", "status": "found", "source": "", "note": ""},
            "2": {"value": "Y", "status": "found", "source": "", "note": ""},
        })})
        answers = InlineBlanksQuery(ai_model=model).result(blanks, "материалы")

        assert model.call_count(Marker.INLINE_BLANKS) == 1
        assert answers["1"]["value"] == "X"


class FakeTemplate:
    """Двойник FormTemplate: markdown задаётся напрямую, файл не нужен.

    Построчному заполнителю от шаблона нужен только текст — писать и
    разбирать .docx ради этого в каждом тесте незачем.
    """

    def __init__(self, markdown: str = "шаблон"):
        self.markdown = markdown
        self.is_word = False
        self.document = None
        self.path = Path("form.docx")


class TestFormTemplate:
    """Шаблон отдаёт оба представления и читает файл не чаще одного раза."""

    def test_markdown_is_parsed_from_the_file(self, template_docx):
        assert "Полное наименование участника" in FormTemplate(str(template_docx)).markdown

    def test_markdown_is_cached(self, template_docx):
        template = FormTemplate(str(template_docx))
        assert template.markdown is template.markdown

    def test_document_is_cached(self, template_docx):
        template = FormTemplate(str(template_docx))
        assert template.document is template.document

    def test_word_template_is_recognised(self, template_docx):
        assert FormTemplate(str(template_docx)).is_word

    def test_non_word_template_is_recognised(self, tmp_path):
        path = tmp_path / "form.xlsx"
        path.write_text("stub", encoding="utf-8")
        assert not FormTemplate(str(path)).is_word

    def test_path_is_exposed_for_the_writer(self, template_docx):
        assert FormTemplate(str(template_docx)).path == Path(str(template_docx))


class TestKnowledgeBaseContext:
    def test_both_sources_are_included(self, requirements_docx, knowledge_dir):
        context = KnowledgeBaseContext(
            tender_info_path=str(requirements_docx),
            knowledge_base_folder=str(knowledge_dir),
        ).build()

        assert "БАЗА ЗНАНИЙ И ЭТАЛОННЫЕ ЗАЯВКИ" in context
        assert "7701234567" in context
        assert "ТРЕБОВАНИЯ ТЕНДЕРА" in context
        assert "Состав заявки участника" in context

    def test_unreadable_tender_info_is_survivable(self, knowledge_dir, tmp_path):
        """Заявку всё равно нужно заполнить тем, что есть в базе знаний."""
        context = KnowledgeBaseContext(
            tender_info_path=str(tmp_path / "нет.docx"),
            knowledge_base_folder=str(knowledge_dir),
        ).build()

        assert "БАЗА ЗНАНИЙ И ЭТАЛОННЫЕ ЗАЯВКИ" in context
        assert "ТРЕБОВАНИЯ ТЕНДЕРА" not in context

    def test_missing_knowledge_base_is_survivable(self, requirements_docx, tmp_path):
        context = KnowledgeBaseContext(
            tender_info_path=str(requirements_docx),
            knowledge_base_folder=str(tmp_path / "нет-базы"),
        ).build()

        assert "ТРЕБОВАНИЯ ТЕНДЕРА" in context
        assert "БАЗА ЗНАНИЙ" not in context

    def test_nothing_available_yields_empty_context(self, tmp_path):
        assert KnowledgeBaseContext(None, None).build() == ""


class TestCompositeTenderForm:
    """Композиция заполнителей: один шаблон, несколько независимых путей."""

    class _StubForm:
        def __init__(self, fields):
            self._fields = fields
            self.calls = []

        def prepare(self, template, context):
            self.calls.append((template, context))
            return list(self._fields)

    def test_results_are_concatenated_in_order(self):
        first = self._StubForm([FilledField(label="А"), FilledField(label="Б")])
        second = self._StubForm([FilledField(label="В")])

        fields = CompositeTenderForm([first, second]).prepare(FakeTemplate(), "к")

        assert [f.label for f in fields] == ["А", "Б", "В"]

    def test_every_form_sees_the_same_template_and_context(self):
        first, second = self._StubForm([]), self._StubForm([])
        template = FakeTemplate()

        CompositeTenderForm([first, second]).prepare(template, "материалы")

        assert first.calls == second.calls == [(template, "материалы")]

    def test_empty_composition(self):
        assert CompositeTenderForm([]).prepare(FakeTemplate(), "к") == []


class TestInlineBlanksForm:
    def test_non_word_template_is_skipped(self):
        """У не-Word шаблона нет объектной модели docx — разбирать нечего."""
        model = ScriptedModel({})
        form = InlineBlanksForm(inline_query=InlineBlanksQuery(ai_model=model))

        assert form.prepare(FakeTemplate(), "контекст") == []
        assert model.calls == []

    def test_template_without_blanks_skips_the_model(self, template_docx):
        model = ScriptedModel({})
        form = InlineBlanksForm(inline_query=InlineBlanksQuery(ai_model=model))

        assert form.prepare(FormTemplate(str(template_docx)), "контекст") == []
        assert model.calls == []

    def test_blanks_become_inline_fields(self, tmp_path):
        path = tmp_path / "form.docx"
        _multi_blank_docx(path)

        model = ScriptedModel({Marker.INLINE_BLANKS: as_json({
            "1": {"value": "САО «ВСК»", "status": "found", "source": "", "note": ""},
            "2": {"value": "Россия", "status": "check", "source": "", "note": ""},
        })})
        form = InlineBlanksForm(inline_query=InlineBlanksQuery(ai_model=model))

        fields = form.prepare(FormTemplate(str(path)), "контекст")

        assert [f.kind for f in fields] == ["inline", "inline"]
        assert [f.anchor for f in fields] == ["1", "2"]
        assert fields[0].value == "САО «ВСК»"
        assert fields[1].status is FieldStatus.CHECK

    def test_blank_without_an_answer_is_marked_missing(self, tmp_path):
        """Модель вернула не все ключи — пропуск не теряется, а помечается."""
        path = tmp_path / "form.docx"
        _multi_blank_docx(path)

        model = ScriptedModel({Marker.INLINE_BLANKS: as_json({
            "1": {"value": "САО «ВСК»", "status": "found", "source": "", "note": ""},
        })})
        form = InlineBlanksForm(inline_query=InlineBlanksQuery(ai_model=model))

        fields = form.prepare(FormTemplate(str(path)), "контекст")

        assert len(fields) == 2
        assert fields[1].status is FieldStatus.MISSING


class TestApplicationOutput:
    def test_output_name_is_derived_from_the_template(self, template_docx, results_dir):
        path = ApplicationOutput(results_path=str(results_dir)).save(
            FormTemplate(str(template_docx)),
            [FilledField(label="ИНН", value="7701234567", status=FieldStatus.FOUND)],
        )

        assert path.name == "form_заполнено.docx"
        assert path.parent == results_dir
        assert path.exists()

    def test_template_is_not_modified_in_place(self, template_docx, results_dir):
        before = template_docx.read_bytes()
        ApplicationOutput(results_path=str(results_dir)).save(
            FormTemplate(str(template_docx)), []
        )
        assert template_docx.read_bytes() == before


class TestAITenderForm:
    def test_extracts_fields_and_fills_them(self, scripted_model):
        form = AITenderForm(tender_query=TenderAIQuery(ai_model=scripted_model))
        fields = form.prepare(FakeTemplate("| ИНН | |"), "контекст")

        assert [f.label for f in fields] == [
            "Полное наименование участника", "ИНН", "Юридический адрес",
            "Контактный телефон", "Руководитель организации",
        ]
        assert fields[1].value == "7701234567"

    def test_statuses_are_mapped(self, scripted_model):
        form = AITenderForm(tender_query=TenderAIQuery(ai_model=scripted_model))
        statuses = {
            f.label: f.status for f in form.prepare(FakeTemplate(), "контекст")
        }

        assert statuses["ИНН"] is FieldStatus.FOUND
        assert statuses["Юридический адрес"] is FieldStatus.CHECK
        assert statuses["Контактный телефон"] is FieldStatus.MISSING

    def test_one_model_call_per_field(self, scripted_model):
        form = AITenderForm(tender_query=TenderAIQuery(ai_model=scripted_model))
        form.prepare(FakeTemplate(), "контекст")

        assert scripted_model.call_count(Marker.FORM_FIELDS) == 1
        assert scripted_model.call_count(Marker.FIELD_VALUE) == 5

    def test_template_reaches_the_extraction_prompt(self, scripted_model):
        form = AITenderForm(tender_query=TenderAIQuery(ai_model=scripted_model))
        form.prepare(FakeTemplate("| Уникальная строка шаблона | |"), "контекст")

        assert "Уникальная строка шаблона" in scripted_model.last_call(Marker.FORM_FIELDS)

    def test_template_without_fields(self):
        model = ScriptedModel({Marker.FORM_FIELDS: as_json({"fields": []})})
        form = AITenderForm(tender_query=TenderAIQuery(ai_model=model))

        assert form.prepare(FakeTemplate(), "контекст") == []
        assert model.call_count(Marker.FIELD_VALUE) == 0

    def test_model_is_required_for_extraction(self):
        class QueryWithoutModel:
            def result(self, field_label, context):
                return {"value": "", "status": "missing", "source": "", "note": ""}

        form = AITenderForm(tender_query=QueryWithoutModel())
        with pytest.raises(RuntimeError, match="модель"):
            form.prepare(FakeTemplate(), "контекст")


class TestTenderApplication:
    """Оркестратор: собрать материалы → отдать заполнителю → сохранить → отчёт.

    Сам он документы не читает и не пишет — это делают FormTemplate,
    KnowledgeBaseContext, TenderForm и ApplicationOutput.
    """

    def _build(self, template, requirements, knowledge, results_dir, model,
               tender_form=None):
        return TenderApplication(
            template=FormTemplate(str(template)),
            context=KnowledgeBaseContext(
                tender_info_path=str(requirements),
                knowledge_base_folder=str(knowledge),
            ),
            tender_form=tender_form or AITenderForm(
                tender_query=TenderAIQuery(ai_model=model)
            ),
            output=ApplicationOutput(results_path=str(results_dir)),
            report=TenderApplicationReport(output_dir=str(results_dir)),
        )

    def test_takes_at_most_five_constructor_parameters(self):
        """Оркестратор не должен обрастать параметрами: всё, что сверх
        композиции зависимостей, — признак утёкшей в него работы."""
        import inspect

        parameters = inspect.signature(TenderApplication.__init__).parameters
        assert len(parameters) - 1 <= 5  # без self

    def test_fills_template_and_saves_it(
        self, template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
    ):
        from docx import Document

        result = self._build(
            template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
        ).result()

        filled = Path(result.filled_path)
        assert filled.exists()
        assert filled.name == "form_заполнено.docx"

        cells = {r.cells[0].text: r.cells[1].text
                 for r in Document(str(filled)).tables[0].rows}
        assert cells["ИНН"] == "7701234567"
        assert cells["Контактный телефон"] == "НЕ НАЙДЕНО"

    def test_context_holds_knowledge_base_and_tender_info(
        self, template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
    ):
        """Часть данных заявки лежит в файле требований — он тоже в контексте."""
        self._build(
            template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
        ).result()

        prompt = scripted_model.last_call(Marker.FIELD_VALUE)
        assert "БАЗА ЗНАНИЙ И ЭТАЛОННЫЕ ЗАЯВКИ" in prompt
        assert "7701234567" in prompt
        assert "ТРЕБОВАНИЯ ТЕНДЕРА" in prompt
        assert "Состав заявки участника" in prompt

    def test_missing_tender_info_is_survivable(
        self, template_docx, knowledge_dir, results_dir, scripted_model, tmp_path
    ):
        application = self._build(
            template_docx, tmp_path / "нет.docx", knowledge_dir,
            results_dir, scripted_model,
        )
        result = application.result()

        assert Path(result.filled_path).exists()
        assert "ТРЕБОВАНИЯ ТЕНДЕРА" not in scripted_model.last_call(Marker.FIELD_VALUE)

    def test_missing_knowledge_base_is_survivable(
        self, template_docx, requirements_docx, results_dir, scripted_model, tmp_path
    ):
        result = self._build(
            template_docx, requirements_docx, tmp_path / "нет-базы",
            results_dir, scripted_model,
        ).result()

        assert Path(result.filled_path).exists()

    def test_report_lists_fields_and_totals(
        self, template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
    ):
        result = self._build(
            template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
        ).result()

        report = Path(result.report_path).read_text(encoding="utf-8")
        assert "Отчёт по заполнению заявки" in report
        assert "Есть (зелёный)" in report
        assert "Проверить (жёлтый)" in report
        assert "Нет (красный)" in report
        assert "- Заполнено: 3" in report
        assert "- Требует проверки: 1" in report
        assert "- Не найдено: 1" in report

    def test_template_is_not_modified_in_place(
        self, template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
    ):
        before = template_docx.read_bytes()
        self._build(
            template_docx, requirements_docx, knowledge_dir, results_dir, scripted_model
        ).result()
        assert template_docx.read_bytes() == before

    def test_inline_blanks_are_filled_end_to_end(
        self, requirements_docx, knowledge_dir, results_dir, tmp_path
    ):
        """Форма с несколькими пропусками в одном абзаце (п. 15.1-стиль):
        построчный механизм такое не различает — каждый пропуск должен
        получить своё, отдельное значение через inline_query."""
        template = tmp_path / "multi_blank_form.docx"
        _multi_blank_docx(template)

        model = CachingScriptedModel({
            # без обычных полей в этом шаблоне — только инлайн-пропуски
            Marker.FORM_FIELDS: as_json({"fields": []}),
            Marker.INLINE_BLANKS: as_json({
                "1": {"value": "САО «ВСК»", "status": "found", "source": "", "note": ""},
                "2": {"value": "Россия", "status": "found", "source": "", "note": ""},
            }),
        })

        result = self._build(
            template, requirements_docx, knowledge_dir, results_dir, model,
            tender_form=CompositeTenderForm([
                AITenderForm(tender_query=TenderAIQuery(ai_model=model)),
                InlineBlanksForm(inline_query=InlineBlanksQuery(ai_model=model)),
            ]),
        ).result()

        statuses = {f.anchor: f.status for f in result.fields if f.kind == "inline"}
        assert len(statuses) == 2
        assert all(status is FieldStatus.FOUND for status in statuses.values())

        # Ответы уходят в колонку справа, текст требования остаётся как был
        row = Document(result.filled_path).tables[0].rows[0]
        assert "САО «ВСК»" in row.cells[1].text
        assert "Россия" in row.cells[1].text
        assert "___" in row.cells[0].text

    def test_form_not_in_the_composition_does_not_run(
        self, requirements_docx, knowledge_dir, results_dir, tmp_path
    ):
        """Заполнитель пропусков не включён в композицию — шаблон с
        пропусками остаётся неразобранным, лишних вызовов модели нет."""
        template = tmp_path / "multi_blank_form.docx"
        _multi_blank_docx(template)

        model = ScriptedModel({Marker.FORM_FIELDS: as_json({"fields": []})})
        result = self._build(
            template, requirements_docx, knowledge_dir, results_dir, model
        ).result()

        assert result.fields == []
        assert model.call_count(Marker.INLINE_BLANKS) == 0
