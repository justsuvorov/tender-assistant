import re
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, List

import pandas as pd
import pdfplumber
from docx import Document
from docx.oxml.ns import qn


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

class MarkdownTableBuilder:
    """Convert tabular data (list of rows) to a GFM markdown table string."""

    @staticmethod
    def from_rows(rows: list[list[str]]) -> str:
        """
        rows[0] is treated as the header row.
        All rows must have the same number of columns (call normalise first).
        """
        if not rows:
            return ""

        col_widths = [3] * len(rows[0])
        for row in rows:
            for i, cell in enumerate(row):
                col_widths[i] = max(col_widths[i], len(cell))

        def fmt_row(values: list[str]) -> str:
            cells = [v.ljust(col_widths[i]) for i, v in enumerate(values)]
            return "| " + " | ".join(cells) + " |"

        separator = "| " + " | ".join("-" * w for w in col_widths) + " |"

        lines = [fmt_row(rows[0]), separator]
        for row in rows[1:]:
            lines.append(fmt_row(row))

        return "\n".join(lines)

    @staticmethod
    def normalise(rows: list[list]) -> list[list[str]]:
        """Pad short rows and stringify every cell."""
        if not rows:
            return []
        max_cols = max(len(r) for r in rows)
        result = []
        for row in rows:
            cells = [str(c).replace("\n", " ").replace("\r", "").strip() if c is not None else "" for c in row]
            cells += [""] * (max_cols - len(cells))
            result.append(cells)
        return result


class WordParagraphConverter:
    """Extract plain text from a Word paragraph XML element, applying heading prefixes."""

    _WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    _HEADING_MAP = {
        "heading1": "#",  "Heading1": "#",
        "heading2": "##", "Heading2": "##",
        "heading3": "###","Heading3": "###",
        "heading4": "####","Heading4": "####",
    }

    def convert(self, element) -> str:
        text = "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()
        if not text:
            return ""

        style_el = element.find(f".//{{{self._WNS}}}pStyle")
        style_val = style_el.get(f"{{{self._WNS}}}val", "") if style_el is not None else ""

        prefix = self._HEADING_MAP.get(style_val, "")
        return f"{prefix} {text}" if prefix else text


class WordTableExtractor:
    """Extract rows from a Word table XML element as list[list[str]]."""

    _WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def extract(self, element) -> list[list[str]]:
        ns = self._WNS
        rows = []
        for tr in element.findall(f"{{{ns}}}tr"):
            cells = []
            for tc in tr.findall(f"{{{ns}}}tc"):
                text = "".join(t.text or "" for t in tc.iter(f"{{{ns}}}t")).strip()
                cells.append(text)
            if cells:
                rows.append(cells)
        return rows


class PdfPageExtractor:
    """Extract text and tables from a single pdfplumber page."""

    def extract(self, page) -> list[str]:
        """Return list of markdown strings (text blocks and tables) for this page."""
        parts = []

        raw_tables = page.extract_tables()
        table_bboxes = [t.bbox for t in page.find_tables()] if raw_tables else []

        for raw_table in raw_tables:
            rows = MarkdownTableBuilder.normalise(raw_table)
            md = MarkdownTableBuilder.from_rows(rows)
            if md:
                parts.append(md)

        text = self._extract_text_outside_tables(page, table_bboxes)
        if text:
            parts.insert(0, text)

        return parts

    @staticmethod
    def _extract_text_outside_tables(page, table_bboxes: list) -> str:
        if not table_bboxes:
            return (page.extract_text() or "").strip()

        remaining = page
        for bbox in table_bboxes:
            try:
                remaining = remaining.outside_bbox(bbox)
            except Exception:
                pass
        return (remaining.extract_text() or "").strip()


# ---------------------------------------------------------------------------
# Parsers
# ---------------------------------------------------------------------------

class Parser(ABC):
    """Base interface for reading documents and converting to markdown."""

    @abstractmethod
    def read_document(self, file_path: str) -> str:
        pass


class Excel(Parser):
    """Read .xlsx/.xls files. Skips empty sheets, converts each sheet to a markdown table."""

    def read_document(self, file_path: str) -> str:
        xl = pd.ExcelFile(file_path)
        sections = []

        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)

            if df.dropna(how="all").empty:
                continue

            df = (
                df.dropna(how="all")
                  .dropna(axis=1, how="all")
                  .reset_index(drop=True)
                  .fillna("")
            )

            rows = [list(df.columns.astype(str))] + df.astype(str).values.tolist()
            rows = MarkdownTableBuilder.normalise(rows)
            md_table = MarkdownTableBuilder.from_rows(rows)

            sections.append(f"## Лист: {sheet_name}\n\n{md_table}")

        return "\n\n".join(sections)


class Word(Parser):
    """Read .docx files. Converts headings, paragraphs, and tables to markdown."""

    def __init__(self):
        self._para_converter = WordParagraphConverter()
        self._table_extractor = WordTableExtractor()

    def read_document(self, file_path: str) -> str:
        doc = Document(file_path)
        parts = []

        for block in doc.element.body:
            tag = block.tag.split("}")[-1]

            if tag == "p":
                md = self._para_converter.convert(block)
                if md:
                    parts.append(md)

            elif tag == "tbl":
                rows = self._table_extractor.extract(block)
                rows = MarkdownTableBuilder.normalise(rows)
                md = MarkdownTableBuilder.from_rows(rows)
                if md:
                    parts.append(md)

        return "\n\n".join(parts)


class PDF(Parser):
    """Read .pdf files. Extracts text and tables page by page."""

    def __init__(self):
        self._page_extractor = PdfPageExtractor()

    def read_document(self, file_path: str) -> str:
        sections = []

        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                parts = self._page_extractor.extract(page)
                if parts:
                    sections.append(
                        f"## Страница {page_num}\n\n" + "\n\n".join(parts)
                    )

        return "\n\n".join(sections)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

class DataParser:
    """Read a file and return its content as a markdown string."""

    _SUPPORTED: dict[str, type[Parser]] = {
        ".xlsx": Excel,
        ".xlsm": Excel,
        ".xls":  Excel,
        ".docx": Word,
        ".doc":  Word,
        ".pdf":  PDF,
    }

    def __init__(self, file_path: str):
        self.file_path = file_path
        self._validate(file_path)
        self.parser = self._build_engine(file_path)

    def origin_data(self, file_path: str=None) -> str:
        """Read document, apply initial cleaning, return markdown string."""
        if file_path is None:
            file_path = self.file_path
        else:
            self._validate(file_path)
        raw = self.parser.read_document(file_path)
        return self._clean(raw)

    @staticmethod
    def _validate(file_path: str) -> None:
        """Fail early with a clear message instead of an SDK-specific error.

        Otherwise a missing path surfaces as PackageNotFoundError from
        python-docx (or a zip error from openpyxl), which callers cannot
        tell apart from a genuine parsing failure.
        """
        if not file_path:
            raise FileNotFoundError("Не указан путь к файлу")

        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Файл не найден: {file_path}")
        if not path.is_file():
            raise FileNotFoundError(f"Путь не является файлом: {file_path}")

    def _build_engine(self, file_path: str) -> Parser:
        ext = Path(file_path).suffix.lower()
        parser_cls = self._SUPPORTED.get(ext)
        if parser_cls is None:
            raise ValueError(
                f"Unsupported file format: '{ext}'. "
                f"Supported: {list(self._SUPPORTED.keys())}"
            )
        return parser_cls()

    @staticmethod
    def _clean(data: str) -> str:
        data = data.strip()
        data = re.sub(r"\n{3,}", "\n\n", data)
        return data


# ---------------------------------------------------------------------------
# Document outline
# ---------------------------------------------------------------------------

class Section:
    """One section of a document: its heading and the text underneath it."""

    def __init__(self, title: str, level: int, body: str = ""):
        self.title = title
        self.level = level
        self.body = body

    @property
    def text(self) -> str:
        """Heading plus body — what gets sent to the LLM as section context."""
        return f"{self.title}\n\n{self.body}".strip()

    def __repr__(self) -> str:
        return f"Section(level={self.level}, title={self.title!r})"


class MarkdownOutline:
    """Split markdown produced by DataParser into sections.

    Two kinds of headings are recognised:

    1. Markdown headings (``#``..``######``) — produced by the Word parser from
       real heading styles and by the PDF/Excel parsers for pages and sheets.
    2. Numbered paragraphs (``5.`` / ``5.1.`` / ``РАЗДЕЛ 3``) — tender
       documents are routinely written without heading styles, so a document
       with no markdown headings at all would otherwise have no outline.

    Heuristic headings are only used when no markdown headings were found,
    so a properly styled document is never polluted by false positives.
    """

    _MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")

    _NUMBERED = re.compile(
        r"^\s*(?:(?:РАЗДЕЛ|Раздел|ГЛАВА|Глава|СТАТЬЯ|Статья|ПРИЛОЖЕНИЕ|Приложение)\s+)?"
        r"(\d+(?:\.\d+)*)\.?\s+(\S.*)$"
    )
    _CAPS = re.compile(r"^[^a-zа-яё]{6,120}$")

    _MAX_HEADING_LEN = 200

    def __init__(self, markdown_text: str):
        self._text = markdown_text or ""
        self._sections = self._parse(self._text)

    @property
    def sections(self) -> list["Section"]:
        return self._sections

    @property
    def has_headings(self) -> bool:
        return bool(self._sections)

    def titles(self) -> list[str]:
        return [s.title for s in self._sections]

    def as_list(self) -> str:
        """Headings as a plain numbered list — the LLM input for section search."""
        return "\n".join(f"{i + 1}. {s.title}" for i, s in enumerate(self._sections))

    def find(self, title: str):
        """Find a section by heading text: exact, then normalised, then partial."""
        if not title:
            return None

        for section in self._sections:
            if section.title == title:
                return section

        target = self._norm(title)
        if not target:
            return None

        for section in self._sections:
            if self._norm(section.title) == target:
                return section

        for section in self._sections:
            norm = self._norm(section.title)
            if norm and (target in norm or norm in target):
                return section

        return None

    def _parse(self, text: str) -> list["Section"]:
        sections = self._parse_markdown(text)
        if sections:
            return sections
        return self._parse_numbered(text)

    def _parse_markdown(self, text: str) -> list["Section"]:
        sections: list[Section] = []
        body: list[str] = []

        for line in text.splitlines():
            match = self._MD_HEADING.match(line.strip())
            if match:
                if sections:
                    sections[-1].body = "\n".join(body).strip()
                body = []
                sections.append(
                    Section(title=match.group(2).strip(), level=len(match.group(1)))
                )
            elif sections:
                body.append(line)

        if sections:
            sections[-1].body = "\n".join(body).strip()

        return sections

    def _parse_numbered(self, text: str) -> list["Section"]:
        sections: list[Section] = []
        body: list[str] = []

        for line in text.splitlines():
            stripped = line.strip()
            level = self._heuristic_level(stripped)

            if level:
                if sections:
                    sections[-1].body = "\n".join(body).strip()
                body = []
                sections.append(Section(title=stripped, level=level))
            elif sections:
                body.append(line)

        if sections:
            sections[-1].body = "\n".join(body).strip()

        return sections

    def _heuristic_level(self, line: str) -> int:
        """Return heading level for a styleless line, or 0 if it is body text."""
        if not line or len(line) > self._MAX_HEADING_LEN or line.startswith("|"):
            return 0

        match = self._NUMBERED.match(line)
        if match:
            tail = match.group(2)
            # A numbered heading is a title, not a sentence: enumeration items
            # end with ";" or "," and sentences contain several full stops.
            if tail.endswith((";", ",")) or tail.count(".") > 1:
                return 0
            return min(match.group(1).count(".") + 1, 6)

        if self._CAPS.match(line) and any(c.isalpha() for c in line):
            return 1

        return 0

    @staticmethod
    def _norm(text: str) -> str:
        text = re.sub(r"^[\s\d.)»«\-–—]+", "", text or "")
        text = re.sub(r"[^\w\s]", " ", text.lower())
        return " ".join(text.split())


# ---------------------------------------------------------------------------
# Inline blanks (несколько пропусков в одном абзаце ячейки таблицы)
# ---------------------------------------------------------------------------

@dataclass
class InlineBlank:
    """Один пропуск внутри абзаца шаблона — не отдельная ячейка, а место
    прямо посреди текста (``_____________ (наименование участника закупки)``).

    Сам текст шаблона НЕ правится: подобранное значение пишется в соседнее
    поле справа (см. reports/writers.py: WordApplicationWriter._fill_inline_blanks),
    откуда пользователь копирует его в нужное место сам. Поэтому здесь
    хранится не только позиция пропуска, но и его положение в таблице:
    ``cell`` — ячейка с пропуском, ``row_cells``/``cell_index`` — строка и
    место в ней, по ним находится ячейка-приёмник ответа.

    ``start``/``end`` — координаты пропуска в тексте абзаца на момент
    сканирования; используются, чтобы отрисовать метки ``[[N]]`` в промпте.
    """

    id: int
    label: str
    paragraph: object
    start: int
    end: int
    cell: object = None
    row_cells: tuple = ()
    cell_index: int = -1


def xml_path(element) -> str:
    """Стабильный ключ идентичности lxml-элемента.

    ``id(element)`` для обёрток python-docx/lxml ненадёжен: сами Python-объекты
    эфемерны (создаются заново при каждом обращении вроде ``row.cells`` или
    ``cell.paragraphs``), и если промежуточный объект не удержан ссылкой,
    сборщик мусора освобождает его адрес — тот же ``id()`` затем достаётся
    СОВСЕМ ДРУГОМУ элементу. На реальном документе (не игрушечном тесте) это
    не редкий случай, а происходит почти на каждой ячейке — проверено эмпирически.
    ``getroottree().getpath(element)`` — путь элемента в дереве, не зависящий
    от времени жизни Python-обёртки, поэтому именно он годится как ключ dict/set.
    """
    return element.getroottree().getpath(element)


_INLINE_BLANK_RUN = re.compile(r"_{3,}|\.{4,}|«_+»")
_INLINE_BLANK_HINT = re.compile(r"^\s*\(([^()]{3,150})\)")

# Одиночный пропуск в ячейке — это уже покрыто механизмом «метка в одной
# ячейке → значение в соседней» (см. application/application.py, TenderAIQuery
# + reports/writers.py, WordApplicationWriter._fill_in_tables). Инлайн-детектор
# нужен только там, где НЕСКОЛЬКО разных по смыслу пропусков делят один
# абзац — той механике такое не по силам в принципе (она даёт одно значение
# на одну ячейку). Порог "2 и больше" — чтобы не дублировать и не
# конфликтовать со старым механизмом на одиночных пропусках.
_MIN_BLANKS_PER_PARAGRAPH = 2


def find_inline_blanks(document: Document) -> List[InlineBlank]:
    """Пропуски (2+ на абзац) во всех ячейках таблиц документа, рекурсивно.

    Работает по реальной объектной модели docx, а не по плоскому markdown —
    только так можно различить несколько пропусков в одном абзаце и
    вернуться потом ровно к нужному месту при записи.
    """
    blanks: List[InlineBlank] = []
    counter = 0

    for cell, row_cells, cell_index in _iter_table_cells(document):
        for paragraph in cell.paragraphs:
            text = paragraph.text
            matches = list(_INLINE_BLANK_RUN.finditer(text))
            if len(matches) < _MIN_BLANKS_PER_PARAGRAPH:
                continue

            for match in matches:
                hint = _INLINE_BLANK_HINT.match(text[match.end():match.end() + 160])
                label = hint.group(1).strip() if hint else _inline_blank_fallback_label(
                    text, match
                )
                counter += 1
                blanks.append(InlineBlank(
                    id=counter, label=label, paragraph=paragraph,
                    start=match.start(), end=match.end(),
                    cell=cell, row_cells=tuple(row_cells), cell_index=cell_index,
                ))

    return blanks


def render_blanks_context(blanks: List[InlineBlank]) -> str:
    """Текст для промпта: абзацы с пропусками, помеченными ``[[N]]``.

    Несколько пропусков одного абзаца показываются одним фрагментом —
    модели нужен полный контекст предложения, а не пропуски по отдельности.
    """
    order: List[str] = []
    grouped: dict = {}
    for blank in blanks:
        key = xml_path(blank.paragraph._p)
        if key not in grouped:
            grouped[key] = []
            order.append(key)
        grouped[key].append(blank)

    fragments = [_annotate_paragraph(grouped[key]) for key in order]
    return "\n\n".join(fragments)


def _annotate_paragraph(group: List[InlineBlank]) -> str:
    text = group[0].paragraph.text
    for blank in sorted(group, key=lambda b: b.start, reverse=True):
        text = text[:blank.start] + f"[[{blank.id}]]" + text[blank.end:]
    return text


def _inline_blank_fallback_label(text: str, match: "re.Match") -> str:
    """Метка, когда рядом с пропуском нет поясняющей скобки."""
    context = text[max(0, match.start() - 60):match.start()].strip()
    return context[-60:] if context else "пропуск без поясняющей подписи"


def _iter_table_cells(document: Document) -> Iterator:
    for table in document.tables:
        yield from _iter_cells_of_table(table)


def _iter_cells_of_table(table) -> Iterator:
    """(ячейка, все ячейки её строки, индекс в строке) — включая вложенные таблицы.

    Позиция в строке нужна, чтобы найти ячейку-приёмник ответа справа от той,
    где стоит пропуск.

    Объединённые по вертикали/горизонтали ячейки: python-docx отдаёт один
    и тот же объект ячейки для каждой строки/колонки, охваченной
    объединением — без дедупликации содержимое объединённой ячейки попало бы
    в результат по разу на каждую охваченную строку. Дедупликация — по
    xml_path(), не по id() (см. его docstring).
    """
    seen = set()
    for row in table.rows:
        cells = row.cells
        for index, cell in enumerate(cells):
            key = xml_path(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            yield cell, cells, index
            for nested_table in cell.tables:
                yield from _iter_cells_of_table(nested_table)
